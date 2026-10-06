"""Monthly MIS - Excel report writer (Accountability Accountants house style).

Reports are never overwritten: a second run for the same company/period is
saved as _v2, _v3 ... beside the first.
"""
from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from skills.monthly_mis.engine import PERIOD_LABELS, PL_LAYOUT

NAVY, TEAL, GREEN, BLUE = "002D49", "174D51", "7DB343", "0A8FD0"
NAVY_TINT, GREEN_TINT, GREY = "DCE6EE", "EAF3DC", "595959"
SEV_FILL = {"High": "F8CBAD", "Medium": "FFE699", "Low": "EAF3DC"}
AED = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
F = "Arial"


def _fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def _safe(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_") or "Company"


def next_version_parts(store, folder: list[str], stem: str) -> list[str]:
    """Never overwrite: first free MIS_..._vN.xlsx in the folder."""
    n = 1
    while store.exists(folder + [f"{stem}_v{n}.xlsx"]):
        n += 1
    return folder + [f"{stem}_v{n}.xlsx"]


class Sheet:
    def __init__(self, wb: Workbook, title: str, heading: str, subtitle: str, widths: list[int], tab=NAVY, first=False):
        self.ws = wb.active if first else wb.create_sheet()
        self.ws.title = title
        self.ws.sheet_properties.tabColor = tab
        self.ws.sheet_view.showGridLines = False
        for i, w in enumerate(widths, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = w
        n = len(widths)
        self.ncol = n
        self.ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
        c = self.ws.cell(1, 1, heading)
        c.font = Font(name=F, size=14, bold=True, color="FFFFFF")
        c.fill = _fill(NAVY)
        c.alignment = Alignment(vertical="center", indent=1)
        for col in range(1, n + 1):
            self.ws.cell(1, col).fill = _fill(NAVY)
        self.ws.row_dimensions[1].height = 28
        s = self.ws.cell(2, 1, subtitle)
        s.font = Font(name=F, size=9, italic=True, color=GREY)
        for col in range(1, n + 1):
            self.ws.cell(3, col).fill = _fill(GREEN)
        self.ws.row_dimensions[3].height = 4
        self.row = 5
        ps = self.ws.page_setup
        ps.orientation = "landscape"
        ps.fitToWidth = 1
        ps.fitToHeight = 0
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True

    def header(self, labels: list[str], freeze=True):
        for i, lab in enumerate(labels, 1):
            c = self.ws.cell(self.row, i, lab)
            c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
            c.fill = _fill(TEAL)
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left" if i == 1 else "center")
        self.ws.row_dimensions[self.row].height = 30
        if freeze:
            self.ws.freeze_panes = self.ws.cell(self.row + 1, 2)
        self.row += 1

    def band(self, text: str):
        for col in range(1, self.ncol + 1):
            self.ws.cell(self.row, col).fill = _fill(GREEN_TINT)
        c = self.ws.cell(self.row, 1, text)
        c.font = Font(name=F, size=10, bold=True, color=NAVY)
        self.row += 1

    def line(self, values: list[Any], fmts: list[str | None] | None = None, style: str = "line", fills: dict[int, str] | None = None):
        bold = style in ("subtotal", "total")
        for i, v in enumerate(values, 1):
            c = self.ws.cell(self.row, i, v)
            c.font = Font(name=F, size=10, bold=bold, color=NAVY if bold else "000000")
            if fmts and i - 1 < len(fmts) and fmts[i - 1]:
                c.number_format = fmts[i - 1]
            if isinstance(v, str) and i > 1:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if style == "subtotal":
                c.fill = _fill(NAVY_TINT)
            if style == "total":
                c.border = Border(top=Side("thin", color=NAVY), bottom=Side("double", color=NAVY))
            if fills and i in fills:
                c.fill = _fill(fills[i])
        self.row += 1

    def note(self, text: str):
        c = self.ws.cell(self.row, 1, text)
        c.font = Font(name=F, size=9, italic=True, color=GREY)
        self.row += 1

    def footer(self, company: str):
        self.ws.oddFooter.left.text = f"Accountability Accountants - {company} - Monthly MIS - &P/&N"
        self.ws.oddFooter.left.font = "Arial"


def _var_pct(a, b):
    return None if not b else (a - b) / abs(b)


def write_report(res: dict[str, Any], commentary: str, commentary_source: str, store,
                 skill_version: str, requested_by: str) -> str:
    """Build the workbook and save it through `store` (local folder or NAS). Returns where it was saved."""
    company, period = res["company"], res["period"] or "Period"
    stem = f"MIS_{_safe(company)}_{_safe(period)}"
    parts = next_version_parts(store, [_safe(company), _safe(period)], stem)
    version_tag = Path(parts[-1]).stem.rsplit("_", 1)[-1]
    sub = (f"{company} | {period} | {res['currency']} | Source: {res['source_file']} | "
           f"Prepared by ACE (Accountability's Chief Examiner) - Monthly MIS Skill v{skill_version} | "
           f"{datetime.now():%d-%b-%Y %H:%M} | Requested by {requested_by} | Report {version_tag} | DRAFT FOR REVIEW")
    pl, k, bs = res["pl"], res["kpis"], res["bs"]
    wb = Workbook()

    # ---------------- Summary ----------------
    s = Sheet(wb, "Summary", f"Monthly MIS - {company} - {period}", sub, [34, 18, 18, 14, 18, 14], first=True)
    s.header(["Headline", "Current Month", "Previous Month", "Change %", "Budget", "vs Budget %"])
    for line in ["Revenue", "Gross Profit", "EBITDA", "Net Profit"]:
        v = pl[line]
        s.line([line, v["cm"], v["pm"], _var_pct(v["cm"], v["pm"]), v["budget"], _var_pct(v["cm"], v["budget"])],
               [None, AED, AED, PCT, AED, PCT], "subtotal" if line != "Net Profit" else "total")
    s.row += 1
    s.header(["KPI (% of revenue)", "Current Month", "Previous Month", "Change (pts)", "Budget", "vs Budget (pts)"], freeze=False)
    for label, v in k.items():
        ch = None if v["cm"] is None or v["pm"] is None else (v["cm"] - v["pm"]) * 100
        cb = None if v["cm"] is None or v["budget"] is None else (v["cm"] - v["budget"]) * 100
        s.line([label, v["cm"], v["pm"], ch, v["budget"], cb], [None, PCT, PCT, '+0.0;-0.0;"-"', PCT, '+0.0;-0.0;"-"'])
    s.row += 1
    s.header(["Checks & exceptions", "Result", "", "", "", ""], freeze=False)
    tbc = res["tb_check"]
    s.line(["Trial balance (current month end)", "Balanced" if tbc["balanced_cm"] else f"OUT BY {tbc['diff_cm']:,.2f}"],
           fills={2: SEV_FILL["Low"] if tbc["balanced_cm"] else SEV_FILL["High"]})
    s.line(["Balance sheet balances", "Yes" if abs(bs["Difference"]["cm"]) <= 1 else f"OUT BY {bs['Difference']['cm']:,.2f}"],
           fills={2: SEV_FILL["Low"] if abs(bs["Difference"]["cm"]) <= 1 else SEV_FILL["High"]})
    for sev in ("High", "Medium", "Low"):
        s.line([f"{sev} exceptions", sum(1 for e in res["exceptions"] if e["severity"] == sev)], fills={2: SEV_FILL[sev]})
    s.line(["Abnormal GL movements", len(res["gl_movements"])])
    s.footer(company)

    # ---------------- P&L ----------------
    p = Sheet(wb, "P&L", f"Profit & Loss - {company} - {period}", sub, [30, 15, 15, 15, 12, 15, 15, 12, 15, 15, 12])
    p.header(["AED", "Current Month", "Previous Month", "Change", "Change %", "Budget", "Variance", "Var %",
              "YTD", "YTD Prior Year", "YTD Change %"])
    for line, kind in PL_LAYOUT:
        v = pl[line]
        p.line([line, v["cm"], v["pm"], v["cm"] - v["pm"], _var_pct(v["cm"], v["pm"]), v["budget"],
                v["cm"] - v["budget"], _var_pct(v["cm"], v["budget"]), v["ytd"], v["ytd_py"], _var_pct(v["ytd"], v["ytd_py"])],
               [None, AED, AED, AED, PCT, AED, AED, PCT, AED, AED, PCT], kind)
    p.row += 1
    p.band("KPIs (% of revenue)")
    for label, v in k.items():
        p.line([label, v["cm"], v["pm"], None, None, v["budget"], None, None, v["ytd"], v["ytd_py"]],
               [None, PCT, PCT, None, None, PCT, None, None, PCT, PCT])
    p.note("Expenses shown as positive amounts; income lines shown as positive amounts. Variance = Actual - Budget.")
    p.footer(company)

    # ---------------- Balance sheet ----------------
    b = Sheet(wb, "Balance Sheet", f"Balance Sheet Summary - {company} - {period}", sub, [32, 18, 18, 18, 12])
    b.header(["AED", "Current Month End", "Previous Month End", "Change", "Change %"])
    bl = bs["lines"]

    def bsl(label, v, style="line"):
        b.line([label, v["cm"], v["pm"], v["cm"] - v["pm"], _var_pct(v["cm"], v["pm"])], [None, AED, AED, AED, PCT], style)

    b.band("Assets")
    for l in ["Non-Current Assets", "Current Assets", "Cash"]:
        bsl(l, bl[l])
    bsl("Total Assets", bs["Total Assets"], "subtotal")
    b.band("Liabilities")
    for l in ["Current Liabilities", "Non-Current Liabilities"]:
        bsl(l, bl[l])
    bsl("Total Liabilities", bs["Total Liabilities"], "subtotal")
    b.band("Equity")
    bsl("Equity (excl. current year profit)", bl["Equity"])
    bsl("Current Year Profit", bl["Current Year Profit"])
    bsl("Total Equity", bs["Total Equity"], "subtotal")
    bsl("Check: Assets - Liabilities - Equity", bs["Difference"], "total")
    b.row += 1
    bsl("Working Capital (CA + Cash - CL)", bs["Working Capital"])
    b.footer(company)

    # ---------------- Branches ----------------
    if res["branches"]:
        br = Sheet(wb, "Branches", f"Branch Performance - {company} - {period}", sub,
                   [26, 15, 15, 11, 15, 11, 15, 10, 10, 15, 15, 10], tab=BLUE)
        br.header(["Branch", "Revenue CM", "Revenue PM", "Change %", "Revenue Budget", "vs Budget %", "Gross Profit CM",
                   "GP % CM", "GP % PM", "EBITDA CM", "EBITDA PM", "EBITDA % CM"])
        for r in res["branches"]:
            fills = {4: SEV_FILL["High"]} if (r["revenue_change_pct"] or 0) <= -0.10 else None
            br.line([r["branch"], r["revenue_cm"], r["revenue_pm"], r["revenue_change_pct"], r["revenue_budget"],
                     r["revenue_vs_budget_pct"], r["gp_cm"], r["gp_pct_cm"], r["gp_pct_pm"], r["ebitda_cm"], r["ebitda_pm"],
                     r["ebitda_pct_cm"]], [None, AED, AED, PCT, AED, PCT, AED, PCT, PCT, AED, AED, PCT], fills=fills)
        br.footer(company)

    # ---------------- GL movements ----------------
    g = Sheet(wb, "GL Movements", f"Abnormal GL Movements - {company} - {period}", sub,
              [12, 34, 22, 15, 15, 15, 11, 15, 28, 30], tab=BLUE)
    g.header(["Code", "Account", "MIS Line", "Current Month", "Previous Month", "Change", "Change %", "Budget", "Flag",
              "Explanation (accounts team)"])
    for r in res["gl_movements"]:
        g.line([r["account_code"], r["account_name"], r["mis_line"], r["cm"], r["pm"], r["change"], r["change_pct"],
                r["budget"], r["flag"], ""], [None, None, None, AED, AED, AED, PCT, AED, None, None], fills={10: "FFFF00"})
    if not res["gl_movements"]:
        g.note("No GL movements above the review thresholds.")
    g.note("Yellow column: to be completed by the accounts team. Amounts presented with income and expenses as positives.")
    g.footer(company)

    # ---------------- Budget variances ----------------
    if res["budget_variances"]:
        bv = Sheet(wb, "Budget Variances", f"Budget Variances - {company} - {period}", sub, [30, 15, 15, 15, 11, 14, 30], tab=BLUE)
        bv.header(["P&L line", "Actual", "Budget", "Variance", "Var %", "Assessment", "Explanation"])
        for r in res["budget_variances"]:
            bv.line([r["line"], r["actual"], r["budget"], r["variance"], r["variance_pct"], r["assessment"], ""],
                    [None, AED, AED, AED, PCT, None, None],
                    fills={6: SEV_FILL["Low"] if r["assessment"] == "Favourable" else SEV_FILL["Medium"], 7: "FFFF00"})
        bv.footer(company)

    # ---------------- Exceptions ----------------
    e = Sheet(wb, "Exceptions", f"Exceptions for Review - {company} - {period}", sub, [10, 16, 52, 60, 15, 26, 14], tab=GREEN)
    e.header(["Severity", "Area", "Item", "Detail / action", "Amount (TB sign)", "Assigned to / response", "Status"])
    for r in res["exceptions"]:
        e.line([r["severity"], r["area"], r["item"], r["detail"], r["amount"], "", "Open"],
               [None, None, None, None, AED, None, None], fills={1: SEV_FILL[r["severity"]], 6: "FFFF00"})
    if not res["exceptions"]:
        e.note("No exceptions identified.")
    e.footer(company)

    # ---------------- Commentary ----------------
    c = Sheet(wb, "Commentary", f"Management Commentary - {company} - {period}", sub, [130], tab=GREEN)
    c.header(["Commentary (draft - review before distribution)"], freeze=False)
    for para in commentary.splitlines():
        cell = c.ws.cell(c.row, 1, para)
        cell.font = Font(name=F, size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        c.row += 1
    c.row += 1
    c.note(f"Commentary source: {commentary_source}. Figures in the commentary must agree to the P&L sheet.")
    c.footer(company)

    # ---------------- Data check ----------------
    d = Sheet(wb, "TB Mapping", f"Trial Balance Mapping - {company} - {period}", sub,
              [12, 36, 24, 16, 10, 15, 15, 15, 15, 15], tab=BLUE)
    d.header(["Code", "Account", "MIS Line", "Branch", "Section", "Current Month", "Previous Month", "Budget", "YTD", "YTD PY"])
    tb = res["_tb"]
    for _, r in tb.iterrows():
        fills = {3: SEV_FILL["High"]} if r["section"] == "UNMAPPED" else None
        d.line([r["account_code"], r["account_name"], r["mis_line"] or "UNMAPPED", r["branch"], r["section"],
                r["cm"], r["pm"], r["budget"], r["ytd"], r["ytd_py"]], [None] * 5 + [AED] * 5, fills=fills)
    d.line(["", "Total (should be nil for YTD)", "", "", "", tb["cm"].sum(), tb["pm"].sum(), tb["budget"].sum(),
            tb["ytd"].sum(), tb["ytd_py"].sum()], [None] * 5 + [AED] * 5, "total")
    d.note("TB sign convention: debit positive, credit negative.")
    d.footer(company)

    buf = io.BytesIO()
    wb.save(buf)
    return store.save(parts, buf.getvalue())

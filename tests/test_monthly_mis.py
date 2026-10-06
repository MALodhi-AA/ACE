from pathlib import Path

import pandas as pd
import pytest
import yaml
from openpyxl import load_workbook

from skills.monthly_mis.engine import MISInputError, analyse, load_input
from skills.monthly_mis.skill import MAPPING_PATH, SKILL_DIR, run

CFG = yaml.safe_load((SKILL_DIR / "skill.yaml").read_text())


@pytest.fixture()
def sample(data_dir):
    return data_dir / "inbox" / "Sample_Restaurant_Sep2026_TB.xlsx"


def test_tb_balances_and_bs_balances(sample):
    res = analyse(load_input(sample, CFG, MAPPING_PATH), CFG)
    assert res["tb_check"]["balanced_cm"]
    assert res["tb_check"]["balanced_pm"]
    assert res["tb_check"]["unmapped"] == 0          # Staff Accommodation mapped via CSV
    assert abs(res["bs"]["Difference"]["cm"]) < 1
    assert abs(res["bs"]["Difference"]["pm"]) < 1


def test_pl_arithmetic(sample):
    inp = load_input(sample, CFG, MAPPING_PATH)
    res = analyse(inp, CFG)
    pl = res["pl"]
    tb = inp.tb[inp.tb["section"] == "PL"]
    # Net profit = -(sum of all P&L TB balances) under Dr+/Cr- convention
    assert pl["Net Profit"]["cm"] == pytest.approx(-tb["cm"].sum(), abs=0.01)
    assert pl["Gross Profit"]["cm"] == pytest.approx(pl["Revenue"]["cm"] - pl["Cost of Sales"]["cm"])
    assert res["kpis"]["EBITDA %"]["cm"] == pytest.approx(pl["EBITDA"]["cm"] / pl["Revenue"]["cm"])


def test_planted_exceptions_found(sample):
    res = analyse(load_input(sample, CFG, MAPPING_PATH), CFG)
    items = " | ".join(e["item"] for e in res["exceptions"])
    assert "Suspense Account" in items
    assert "Expenses to be allocated" in items
    assert "Due from Related Party" in items
    assert "Dubai Mall" in items and "revenue down" in items.lower()
    assert "Food Cost %" in items
    flagged = {g["account_name"] for g in res["gl_movements"]}
    assert {"Repairs and Maintenance", "Expenses to be allocated"} <= flagged


def test_unbalanced_and_unmapped(tmp_path, sample):
    df = pd.read_excel(sample, sheet_name="TB")
    df.loc[len(df)] = ["9999", "Mystery account", "", "", 5000, 0, 0, 5000, 0]
    bad = tmp_path / "bad.xlsx"
    with pd.ExcelWriter(bad) as xw:
        pd.read_excel(sample, sheet_name="Info", header=None).to_excel(xw, sheet_name="Info", header=False, index=False)
        df.to_excel(xw, sheet_name="TB", index=False)
    res = analyse(load_input(bad, CFG, MAPPING_PATH), CFG)
    assert not res["tb_check"]["balanced_cm"]
    sev = [(e["severity"], e["item"]) for e in res["exceptions"]]
    assert ("High", "Unmapped account: 9999 Mystery account") in sev


def test_run_writes_versioned_report(data_dir):
    r1 = run("sample_restaurant", requested_by="test", use_ai=False)
    r2 = run("Sample_Restaurant_Sep2026_TB.xlsx", requested_by="test", use_ai=False)
    assert r1.ok and r2.ok
    p1, p2 = Path(r1.report_path), Path(r2.report_path)
    v1 = int(p1.stem.rsplit("_v", 1)[1])
    v2 = int(p2.stem.rsplit("_v", 1)[1])
    assert v2 == v1 + 1 and p1.exists()      # never overwritten
    wb = load_workbook(p1)
    assert {"Summary", "P&L", "Balance Sheet", "Branches", "GL Movements", "Exceptions", "Commentary", "TB Mapping"} <= set(wb.sheetnames)
    assert "MIS completed" in r1.chat_summary


def test_missing_file_reported(data_dir):
    r = run("does_not_exist.xlsx", use_ai=False)
    assert not r.ok and "No file" in r.chat_summary


def test_mis_from_client_source(data_dir):
    r = run("Clients/Mara/2026/09 Sep/mara tb", requested_by="test", use_ai=False)
    assert r.ok, r.chat_summary
    assert r.details["source"] == "Clients/Mara/2026/09 Sep/Mara TB Sep 2026.xlsx"

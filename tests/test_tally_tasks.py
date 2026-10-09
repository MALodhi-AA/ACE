"""v0.8.0: client register, Tally check types, Tally tasks and the chat commands."""
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

from app import tally_tasks as tt
from app.storage import LocalStore
from app.tally_checks import Ctx, period, run_check
from app.tally_register import REGISTER, draft, load
from integrations.tally.client import Ledger
from tests.test_assistant import DM_MA, MA, last_to, office  # noqa: F401 - fixture

TODAY = date(2026, 10, 7)
CO = "Mara Lounge Al Wasl (E-1)"
CO2 = "Volt Electric FZE"


class FakeTally:
    """Two companies; balances depend on the 'to' date (like Tally with typed dates)."""

    def __init__(self):
        self.calls = []
        self.future = []                 # vouchers dated after today
        self.suspense = -1500.0          # Dr
        self.postings = []

    def companies(self):
        return [{"name": CO, "starting_from": date(2019, 1, 1), "books_from": date(2019, 1, 1)},
                {"name": CO2, "starting_from": date(2024, 4, 1), "books_from": date(2024, 4, 1)}]

    def groups(self, company):
        return {"Current Assets": "", "Bank Accounts": "Current Assets", "Cash-in-Hand": "Current Assets",
                "Suspense A/c": "", "Sales Accounts": "", "Local Banks": "Bank Accounts"}

    def ledgers(self, company, frm=None, to=None):
        self.calls.append(("ledgers", company, frm, to))
        sales_open, sales_close = 0.0, 0.0
        if frm and to:                   # sales movement: Sep 100,000; Aug 60,000
            sales_close = {9: 100000.0, 8: 60000.0}.get(to.month, 50000.0)
        return [Ledger("Emirates NBD", "Local Banks", -10.0, -5000.0),
                Ledger("Cash", "Cash-in-Hand", 0.0, 250.0 if company == CO else -100.0),   # CO: cash in credit
                Ledger("Suspense A/c", "Suspense A/c", 0.0, self.suspense if company == CO else 0.0),
                Ledger("Sales", "Sales Accounts", sales_open, sales_close)]

    def voucher_dates(self, company, frm=None, to=None, strict=True, extra=None):
        self.calls.append(("dates", company, frm, to, extra))
        return [date(2026, 10, 5)] if company == CO else [date(2026, 8, 1)]

    def last_voucher_date(self, company, frm=None, to=None):
        d = self.voucher_dates(company, frm, to)
        return d[-1] if d else None

    def vouchers(self, company, frm, to):
        if frm > TODAY:
            return [v for v in self.future if company == CO]
        return [v for v in self.postings if frm <= v["date"] <= to] if company == CO else []

    def trial_balance(self, company, frm=None, to=None):
        return [("Capital Account", 5000.0), ("Current Assets", -5000.0)]


@pytest.fixture
def ft():
    return FakeTally()


def test_periods():
    assert period("last_7_days", TODAY) == (date(2026, 10, 1), TODAY)
    assert period("last_month", TODAY) == (date(2026, 9, 1), date(2026, 9, 30))
    assert period("last_quarter", TODAY) == (date(2026, 7, 1), date(2026, 9, 30))
    assert period("this_week", TODAY) == (date(2026, 10, 5), TODAY)
    assert period("2026-01-01..2026-03-31", TODAY) == (date(2026, 1, 1), date(2026, 3, 31))
    assert period("this_year", TODAY, date(2026, 4, 1)) == (date(2026, 4, 1), TODAY)


def test_checks_find_problems(ft, tmp_path):
    ctx = Ctx(ft, today=TODAY, store=LocalStore(tmp_path))
    ft.future = [{"date": date(2030, 9, 15), "number": "45", "type": "Sales", "narration": "", "party": "", "lines": []}]
    r = run_check(ctx, "future_entries", CO, {})
    assert r.status == "fail" and r.summary == "1 likely error (dated up to 15 Sep 2030): Sales 1"
    assert run_check(ctx, "future_entries", CO2, {}).status == "ok"

    assert run_check(ctx, "last_entry", CO, {"max_days": 7}).status == "ok"
    r = run_check(ctx, "last_entry", CO2, {"max_days": 7, "voucher_type": "sales"})
    assert r.status == "fail" and "last sales entry 01 Aug 2026" in r.summary
    assert ft.calls[-1][4] == {"ACEType": "$$IsSales:$VoucherTypeName"}

    r = run_check(ctx, "ledger_balance", CO, {"ledger": "suspense", "rule": "zero"})
    assert r.status == "fail" and r.details == ["Suspense A/c: 1,500.00 Dr"]
    assert ("ledgers", CO, date(2026, 1, 1), TODAY) in ft.calls               # FY start from Tally (Jan)
    assert run_check(ctx, "ledger_balance", CO2, {"ledger": "suspense", "rule": "zero"}).status == "ok"
    r = run_check(ctx, "ledger_balance", CO, {"ledger": "cash", "rule": "no_credit"})
    assert r.status == "fail" and "Cash: 250.00 Cr" in r.details
    r = run_check(ctx, "ledger_balance", CO, {"group": "Bank Accounts", "rule": "max", "amount": 1000})
    assert r.details == ["Emirates NBD: 5,000.00 Dr"]                        # sub-group Local Banks included
    assert run_check(ctx, "ledger_balance", CO, {"ledger": "Nope"}).status == "error"

    ft.postings = [{"date": date(2026, 10, 6), "number": "J-12", "type": "Journal", "narration": "unknown receipt",
                    "party": "", "lines": [("Suspense A/c", -1500.0), ("Emirates NBD", 1500.0)]}]
    r = run_check(ctx, "ledger_postings", CO, {"ledger": "suspense", "period": "last_7_days"})
    assert r.status == "fail" and r.details[0] == "06 Oct 2026 Journal J-12 - 1,500.00 - Suspense A/c: 1,500.00 Dr - unknown receipt"

    r = run_check(ctx, "compare", CO, {"ledger": "Sales", "period": "last_month", "threshold_pct": 30})
    assert r.status == "fail" and "+67%" in r.summary                         # 100,000 vs 60,000

    r = run_check(ctx, "report", CO, {"kind": "trial_balance"})
    assert r.status == "info" and r.file.endswith("Mara Lounge Al Wasl (E-1) - Trial Balance 2026-10-07.xlsx")
    ws = load_workbook(r.file).active
    assert ws["A1"].value.startswith("Trial Balance - Mara Lounge Al Wasl (E-1)")            # branded title bar
    assert ws["A1"].fill.start_color.rgb.endswith("002D49") and ws["A5"].fill.start_color.rgb.endswith("174D51")
    assert ws["A5"].value == "Particulars" and ws["A1"].font.name == "Arial" and not ws.sheet_view.showGridLines
    assert ws["A6"].value == "Capital Account" and ws["C6"].value == 5000.0 and ws["B7"].value == 5000.0
    assert ws["A8"].value == "Total" and ws["B8"].value == ws["C8"].value == 5000.0
    assert ws["B8"].border.bottom.style == "double"
    assert ws.oddFooter.left.text.startswith("Accountability Accountants - Mara Lounge Al Wasl (E-1) - Trial Balance")


def test_register_draft_and_load(ft, tmp_path):
    store = LocalStore(tmp_path)
    where, n = draft(ft, store=store, today=TODAY)
    assert n == 2 and where.endswith("Client Register (draft) 2026-10-07.xlsx")
    ws = load_workbook(where).worksheets[0]
    assert ws["A5"].value == "Tally company" and ws["A1"].fill.start_color.rgb.endswith("002D49")
    row = [c.value for c in ws[6]]
    assert ws["D6"].fill.start_color.rgb.endswith("FFFF00")                 # staff to be filled in
    assert row[0] == CO and row[1] == "Mara Lounge Al Wasl" and row[7] == "1-Jan"
    assert row[9] == "Emirates NBD" and row[10] == "Cash" and row[11] == "Suspense A/c" and row[12] == "05 Oct 2026"
    assert draft(ft, store=store, today=TODAY)[0].endswith("2026-10-07_v2.xlsx")       # never overwrites

    wb = load_workbook(where)
    wb.worksheets[0]["B6"] = "Al Wasl"
    wb.worksheets[0]["D6"] = "Aiman"
    wb.worksheets[0]["C7"] = "N"
    wb.save(tmp_path / REGISTER)
    reg = load(store)
    assert reg[CO].short == "Al Wasl" and reg[CO].staff == "Aiman" and not reg[CO2].include
    names = [c["name"] for c in ft.companies()]
    assert tt.select_companies("all", names, reg) == ([CO], [])
    assert tt.select_companies({"staff": "aiman"}, names, reg) == ([CO], [])
    assert tt.select_companies(["al wasl", "Zeta"], names, reg) == ([CO], ["Zeta"])


def test_schedules():
    t = datetime(2026, 10, 7, 11, 0, tzinfo=tt.tz())                          # Wednesday
    assert tt.parse_schedule("Weekly Monday 9:30") == "weekly mon 09:30"
    assert tt.next_run("weekly mon 09:30", t) == datetime(2026, 10, 12, 9, 30, tzinfo=tt.tz())
    assert tt.next_run("daily 10:00", t) == datetime(2026, 10, 8, 10, 0, tzinfo=tt.tz())
    assert tt.next_run("weekdays 12:00", t) == datetime(2026, 10, 7, 12, 0, tzinfo=tt.tz())
    assert tt.next_run("monthly 5 08:00", t) == datetime(2026, 11, 5, 8, 0, tzinfo=tt.tz())
    assert tt.next_run("manual", t) is None
    with pytest.raises(ValueError):
        tt.parse_schedule("every full moon")


def test_chat_create_list_run_and_schedule(office, ft, tmp_path):  # noqa: F811
    chat, w, replies, clock, say = office
    a = w.assistant
    a.tally, a.tally_store = ft, LocalStore(tmp_path / "ace-tally")
    say(MA, "hi")
    replies["next"] = lambda s, u: {"name": "Suspense must be zero", "check": "ledger_balance",
                                    "params": {"ledger": "suspense", "rule": "zero", "junk": 1},
                                    "companies": "all", "schedule": "weekly mon 10:00", "report": "exceptions"}
    say(MA, "tally task: every Monday at 10 check all companies, suspense must be zero")
    c = last_to(chat, DM_MA)
    assert c.startswith("New Tally task:") and "Check: ledger_balance (ledger=suspense, rule=zero)" in c
    assert "When: every Mon at 10:00" in c and "Reply YES to save it" in c
    say(MA, "yes")
    assert last_to(chat, DM_MA).startswith("Saved as Tally task 1.")
    t = a.tally_tasks.get(1)
    assert t["next_run"].startswith("2026-10-12T10:00") and "junk" not in t["params"]

    say(MA, "tally tasks")
    assert "1. Suspense must be zero - ledger_balance" in last_to(chat, DM_MA)
    say(MA, "tally run 1")
    out = last_to(chat, DM_MA)
    assert out.startswith("Tally check 1 'Suspense must be zero' - 2 companies: 1 clear, 1 need attention")
    assert "- Mara Lounge Al Wasl (E-1): suspense ledger should be zero (as at 07 Oct 2026)" in out
    assert "    Suspense A/c: 1,500.00 Dr" in out
    say(MA, "tally run 1 for volt")
    assert last_to(chat, DM_MA).startswith("Tally check 1 'Suspense must be zero' - 1 company: 1 clear")

    clock.t = datetime(2026, 10, 12, 10, 1, tzinfo=clock.t.tzinfo)           # Monday 10:01: scheduled run
    before = len(chat.sent)
    a.tick(force=True)
    w._pool.shutdown(wait=True)
    w._pool = type(w._pool)(max_workers=4)
    assert any(c == DM_MA and t.startswith("Tally check 1") for c, _, t in chat.sent[before:])
    assert a.tally_tasks.get(1)["next_run"].startswith("2026-10-19T10:00")
    say(MA, "tally show 1")
    assert "Recent findings:" in last_to(chat, DM_MA)
    say(MA, "tally pause 1")
    assert a.tally_tasks.get(1)["enabled"] == 0
    say(MA, "tally delete 1")
    assert a.tally_tasks.get(1) is None


def test_task_draft_with_problem_is_not_saved(office, ft):  # noqa: F811
    chat, w, replies, clock, say = office
    w.assistant.tally = ft
    say(MA, "hi")
    replies["next"] = lambda s, u: {"name": "x", "check": "ledger_balance", "params": {}, "schedule": "sometimes"}
    say(MA, "tally task: check balances")
    c = last_to(chat, DM_MA)
    assert "Please fix:" in c and "which ledger or group?" in c and "can't read the schedule" in c
    say(MA, "yes")
    assert "still has open points" in last_to(chat, DM_MA)
    replies["next"] = lambda s, u: {"name": "Cash", "check": "ledger_balance",
                                    "params": {"ledger": "cash", "rule": "no_credit"}, "schedule": "daily 09:00"}
    say(MA, "cash ledger, must not be in credit, daily at 9")
    assert "Reply YES to save it" in last_to(chat, DM_MA)
    say(MA, "no")
    assert last_to(chat, DM_MA).startswith("Dropped") and w.assistant.tally_tasks.all() == []


def test_tally_question_through_the_brain(office, ft, monkeypatch):  # noqa: F811
    chat, w, replies, clock, say = office
    a = w.assistant
    a.tally = ft
    a._fast_brain = lambda system, user: {"tool": "tally_balances", "args": {"company": "al wasl", "ledger": "cash"},
                                          "complex": False}
    w.run_later = lambda f: f()
    say(MA, "what is the cash balance at al wasl?")
    out = last_to(chat, DM_MA)
    assert out.startswith("Mara Lounge Al Wasl (E-1) - cash ledger as at 07 Oct 2026 (total 250.00 Cr)")
    say(MA, "tally companies")
    assert "2 companies open in Tally (no client register yet)" in last_to(chat, DM_MA)


def test_tally_not_connected(office):  # noqa: F811
    chat, w, replies, clock, say = office
    say(MA, "tally tasks")
    assert last_to(chat, DM_MA) == "Tally is not connected (set TALLY_URL in .env)."


def test_logo_goes_top_left_when_available(tmp_path):
    import io

    from openpyxl import Workbook
    from PIL import Image

    from app.brand import Sheet
    buf = io.BytesIO()
    Image.new("RGB", (200, 60), "white").save(buf, format="PNG")
    wb = Workbook()
    sh = Sheet(wb, "S", "Title", "sub", [20, 20], first=True, logo_bytes=buf.getvalue())
    sh.header(["a", "b"])
    assert len(sh.ws._images) == 1 and sh.ws["A2"].value == "Title" and sh.ws["A6"].value == "a"


def test_closed_companies_are_left_out_of_all_companies(ft, tmp_path):
    closed = "Z (Closed) Mara JLT Branch (E-1)"
    names = [CO, closed]
    assert tt.select_companies("all", names, {}) == ([CO], [])
    assert tt.select_companies([closed], names, {}) == ([closed], [])         # still checkable by name
    ft.companies = lambda: [{"name": closed, "starting_from": date(2023, 1, 1), "books_from": None}]
    where, _ = draft(ft, store=LocalStore(tmp_path), today=TODAY)
    assert load_workbook(where).worksheets[0]["C6"].value == "N"


def test_same_error_in_many_companies_is_one_line(ft, tmp_path):
    store = TaskStoreForTest(tmp_path)
    tasks = tt.TallyTasks(store)
    task = tasks.add({"name": "Bad", "check": "ledger_balance", "params": {"ledger": "Nope"}, "schedule": "manual"})
    names = [f"Co {i}" for i in range(6)]
    ft.companies = lambda: [{"name": n, "starting_from": date(2026, 1, 1), "books_from": None} for n in names]
    out = tt.run_task(task, ft, tasks, register={}, today=TODAY)
    assert "- 6 companies could not be checked - no ledger 'Nope' in this company: Co 0, Co 1, Co 2, Co 3, Co 4 ..." in out


def TaskStoreForTest(tmp_path):
    from app.tasks import TaskStore
    return TaskStore(tmp_path / "t.db")


def test_future_entries_errors_vs_planned_entries_and_excel_list(ft, tmp_path):
    from app.tasks import TaskStore
    tasks = tt.TallyTasks(TaskStore(tmp_path / "t.db"))
    task = tasks.add({"name": "Future", "check": "future_entries", "params": {}, "schedule": "manual"})
    jv = [{"date": date(2026, 10 + i % 3, 28), "number": str(500 + i), "type": "Journal", "party": "",
           "narration": "Being rent   amortisation for the month" if i % 2 else "",
           "lines": [("Rent", -1000.0), ("Prepaid rent", 1000.0)]} for i in range(12)]
    pdc = [{"date": date(2027, 1, 15), "number": "1", "type": "Chq-Corres.", "party": "City Center", "narration": "",
            "lines": [("City Center", -450044.0), ("Bank", 450044.0)]}]
    inv = {"date": date(2026, 11, 30), "number": "CI-2025-01969", "type": "Pur-Credit-Jun", "party": "SANCO",
           "narration": "Inv dated 30.11.2025", "lines": [("SANCO", 1050.0), ("Purchases", -1000.0),
                                                           ("Input VAT", -50.0)]}
    ft.future = jv + pdc + [inv]
    out = tt.run_task(task, ft, tasks, register={}, today=TODAY, store=LocalStore(tmp_path / "t"))
    lines = out.split("\n")
    assert "- Mara Lounge Al Wasl (E-1): 1 likely error (dated up to 30 Nov 2026): Pur-Credit-Jun 1" in lines
    assert "    30 Nov 2026 Pur-Credit-Jun CI-2025-01969 - SANCO - 1,050.00 - Inv dated 30.11.2025" in lines
    assert "    Post-dated cheques: 1 (15 Jan 2027) - not counted" in lines
    assert "    Journals in advance: 12 (28 Oct 2026 - 28 Dec 2026), 6 without narration - not counted" in lines
    assert "Journal 500" not in out                                   # planned entries are not listed one by one
    path = out.split("Full list with narrations (every entry, grouped): ")[1].strip()
    ws = load_workbook(path).active
    head = [c.value for c in ws[5]]
    assert head == ["Company", "Date", "Type", "Number", "Party", "Amount", "Group", "Narration"]
    assert ws["G7"].value == "Likely error" and ws.max_row == 6 + 14

    task2 = tasks.add({"name": "All", "check": "future_entries", "params": {"include_journals": "yes"},
                       "schedule": "manual"})
    out = tt.run_task(task2, ft, tasks, register={}, today=TODAY, store=LocalStore(tmp_path / "t"))
    assert "13 entries dated in the future (dated up to 28 Dec 2026): Journal 12, Pur-Credit-Jun 1" in out

    ft.future = jv
    r = run_check(Ctx(ft, today=TODAY), "future_entries", CO, {})
    assert r.status == "ok" and r.summary.startswith("no likely errors; Journals in advance: 12")


def test_journals_without_narration(ft):
    ft.postings = [{"date": date(2026, 9, 30), "number": "J-1", "type": "Journal", "party": "", "narration": "",
                    "lines": [("Rent", -10.0), ("Bank", 10.0)]},
                   {"date": date(2026, 9, 30), "number": "J-2", "type": "Journal", "party": "", "narration": "accrual",
                    "lines": [("Rent", -10.0), ("Bank", 10.0)]},
                   {"date": date(2026, 9, 12), "number": "P-9", "type": "Payment", "party": "", "narration": "",
                    "lines": [("Rent", -10.0), ("Bank", 10.0)]}]
    r = run_check(Ctx(ft, today=TODAY), "journals_without_narration", CO, {"period": "last_month"})
    assert r.status == "fail" and r.summary == "1 journal without narration 01 Sep - 30 Sep 2026"
    r = run_check(Ctx(ft, today=TODAY), "journals_without_narration", CO, {"voucher_type": "any"})
    assert r.summary.startswith("2 vouchers without narration")

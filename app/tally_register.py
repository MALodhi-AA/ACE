"""Client register for Tally (v0.8.5): groups, entities, units and inter-company ledgers.

Structure (agreed with Sir Muhammad Ali):
  Group  ->  Entity (E-1, E-2 ... the legal / tax person)  ->  Unit (one Tally company: outlet,
  branch, kitchen ...). Units with the same entity code are ONE entity for tax, whatever the brand.
  A unit with no entity code is a stand-alone entity. Transactions inside an entity (or, later,
  inside a VAT / CT tax group) have no tax effect.

Two sheets, drafted by ACE from Tally ("draft client register"), completed by the team and saved
as ACE/tally/Client Register.xlsx:
  Units          one row per Tally company: group, entity, tax details, staff, ledgers ...
                 (entity details are repeated on each unit; ACE checks they agree)
  Inter-company  which ledger in a unit stands for which other unit (Branch / Divisions,
                 Sundry Debtors / Creditors), confirmed Y by the team.

ACE fills what it can find itself; yellow cells are for the team. Drafts never overwrite.
"""
from __future__ import annotations

import difflib
import io
import logging
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, timedelta

from app.config import settings

log = logging.getLogger(__name__)

FOLDER = ["tally"]
REGISTER = "Client Register.xlsx"
STANDALONE = "stand-alone?"
UNIT_COLS = ["Tally company", "Short name", "Include (Y/N)", "Group", "Entity code", "Entity legal name", "TRN",
             "VAT period", "VAT quarter ends", "CT TRN", "FY start", "VAT group", "CT tax group", "Responsible staff",
             "Manager", "Bank ledgers", "Cash ledgers", "Suspense ledgers", "Last entry (when drafted)", "Notes"]
UNIT_WIDTHS = [38, 22, 9, 16, 10, 30, 18, 11, 16, 18, 9, 14, 14, 18, 18, 36, 26, 26, 15, 30]
IC_COLS = ["Tally company", "Ledger", "Ledger group", "Counterparty (Tally company)", "Match", "Relationship",
           "Confirmed (Y/N)", "Notes"]
IC_WIDTHS = [38, 36, 22, 38, 10, 18, 11, 30]
ENTITY_FIELDS = {"entity_name": "Entity legal name", "trn": "TRN", "vat_period": "VAT period",
                 "vat_quarter_ends": "VAT quarter ends", "ct_trn": "CT TRN", "fy_start": "FY start",
                 "vat_group": "VAT group", "ct_group": "CT tax group", "group": "Group"}
IC_GROUPS = ("Branch / Divisions", "Sundry Debtors", "Sundry Creditors")
HELP = [
    ("Units sheet", "One row per Tally company (unit). Yellow = to be completed or confirmed by the team."),
    ("Tally company", "Exact company name in Tally - do not change."),
    ("Short name", "What you call the unit in chat (e.g. Mara Al Wasl). ACE accepts it in instructions."),
    ("Include (Y/N)", "N = left out of 'all companies' checks (closed units are drafted as N)."),
    ("Group", "Commercial group (e.g. Mara). ACE suggests it from the names - please confirm."),
    ("Entity code", "E-1, E-2 ... the legal / tax entity. Units with the same code are ONE entity, whatever the brand. "
     "ACE takes it from the name. 'stand-alone?' = no code found: leave the cell empty if the unit is its own "
     "entity, or enter the entity code it belongs to."),
    ("Entity legal name, TRN, VAT period, VAT quarter ends, CT TRN, FY start",
     "Entity details - repeat them on every unit of the entity (ACE warns when units of one entity disagree). "
     "VAT period: Monthly / Quarterly / None. Quarter ends e.g. Mar Jun Sep Dec. FY start e.g. 1-Jan."),
    ("VAT group, CT tax group", "Leave empty. When entities form a VAT group or a CT tax group, put the same group "
     "name on all their units - ACE then treats transactions between them as inside the tax group."),
    ("Responsible staff, Manager", "Chat names - used when you ask ACE to take a finding up."),
    ("Bank / Cash / Suspense ledgers", "Found by ACE from Tally groups and names; separate several with ';'."),
    ("Inter-company sheet", "Ledgers in each unit that stand for another unit (found under Branch / Divisions and "
     "Sundry Debtors / Creditors). ACE suggests the counterparty; Match shows how sure it is."),
    ("Confirmed (Y/N)", "Y = ACE uses this link in inter-company checks. Exact name matches are pre-set to Y; check "
     "the others. A Branch / Divisions ledger with no counterparty: enter the Tally company, or write 'outside "
     "Tally' in Notes."),
    ("Relationship", "Worked out by ACE from the Units sheet: same entity / same VAT group / same group / "
     "different group. No need to fill."),
]


@dataclass
class Client:
    company: str
    short: str = ""
    include: bool = True
    group: str = ""
    entity: str = ""                 # E-n, or "" = stand-alone
    entity_name: str = ""
    trn: str = ""
    vat_period: str = ""
    vat_quarter_ends: str = ""
    ct_trn: str = ""
    fy_start: str = ""
    vat_group: str = ""
    ct_group: str = ""
    staff: str = ""
    manager: str = ""
    roles: dict[str, list[str]] = field(default_factory=dict)      # bank / cash / suspense -> ledgers
    notes: str = ""

    @property
    def label(self) -> str:
        return self.short or self.company

    @property
    def entity_key(self) -> str:
        return self.entity or f"unit:{self.company}"


@dataclass
class Link:
    company: str
    ledger: str
    ledger_group: str
    counterparty: str
    match: str = ""
    confirmed: bool = False
    notes: str = ""


def _store():
    from app.nas import nas
    from app.storage import LocalStore, NasStore
    if nas.configured:
        return NasStore(settings.ace_share, FOLDER)
    return LocalStore(settings.data_dir / "tally")


def _read(store, name: str) -> bytes | None:
    from app.storage import LocalStore
    try:
        if isinstance(store, LocalStore):
            p = store.root / name
            return p.read_bytes() if p.exists() else None
        from app.nas import nas
        if not store.exists([name]):
            return None
        return nas.read_bytes(store.share, store.base + [name])
    except Exception as exc:  # noqa: BLE001
        log.warning("client register: %s", exc)
        return None


def _split(v) -> list[str]:
    return [x.strip() for x in re.split(r"[;\n]", str(v or "")) if x.strip()]


def _table(ws, first: str) -> tuple[list[str], list[tuple]]:
    """(headers, rows) of a sheet whose header row starts with `first` (title block above is skipped)."""
    rows = list(ws.iter_rows(values_only=True))
    start = next((i for i, r in enumerate(rows) if r and str(r[0] or "").strip().lower() == first.lower()), None)
    if start is None:
        return [], []
    body = [r for r in rows[start + 1:] if r and any(v not in (None, "") for v in r[1:])]   # skip notes / blank rows
    return [str(h or "").strip().lower() for h in rows[start]], body


def _col(head, r, name) -> str:
    try:
        v = r[head.index(name.lower())]
    except (ValueError, IndexError):
        return ""
    return "" if v is None else str(v).strip()


def entity_code(text: str) -> str:
    """'E-1' from '... (E-1)' / 'E1' / 'e - 1'; '' if none."""
    m = re.search(r"\bE\s*-?\s*(\d+)\b", text or "", re.IGNORECASE)
    return f"E-{int(m.group(1))}" if m else ""


def _workbook(store):
    data = _read(store, REGISTER)
    if not data:
        return None
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(data), read_only=True, data_only=True)


def load(store=None) -> dict[str, Client]:
    """Units of the filled-in register, by Tally company name ({} if there is none yet)."""
    wb = _workbook(store or _store())
    if wb is None:
        return {}
    ws = next((s for s in wb.worksheets if _table(s, "Tally company")[0] and "ledger" not in _table(s, "Tally company")[0]),
              wb.worksheets[0])
    head, rows = _table(ws, "Tally company")
    out = {}
    for r in rows:
        company = _col(head, r, "Tally company")
        if not company:
            continue
        ent = _col(head, r, "Entity code")
        out[company] = Client(
            company=company, short=_col(head, r, "Short name"),
            include=_col(head, r, "Include (Y/N)").upper() not in ("N", "NO"),
            group=_col(head, r, "Group"),
            entity="" if ent.lower().startswith("stand") else (entity_code(ent) or ent),
            entity_name=_col(head, r, "Entity legal name"), trn=_col(head, r, "TRN"),
            vat_period=_col(head, r, "VAT period"), vat_quarter_ends=_col(head, r, "VAT quarter ends"),
            ct_trn=_col(head, r, "CT TRN"), fy_start=_col(head, r, "FY start"),
            vat_group=_col(head, r, "VAT group"), ct_group=_col(head, r, "CT tax group"),
            staff=_col(head, r, "Responsible staff"), manager=_col(head, r, "Manager"), notes=_col(head, r, "Notes"),
            roles={"bank": _split(_col(head, r, "Bank ledgers")), "cash": _split(_col(head, r, "Cash ledgers")),
                   "suspense": _split(_col(head, r, "Suspense ledgers"))})
    return out


def load_links(store=None, confirmed_only: bool = True) -> list[Link]:
    """Inter-company links from the register."""
    wb = _workbook(store or _store())
    if wb is None:
        return []
    out = []
    for ws in wb.worksheets:
        head, rows = _table(ws, "Tally company")
        if "ledger" not in head:
            continue
        for r in rows:
            company, ledger = _col(head, r, "Tally company"), _col(head, r, "Ledger")
            if not company or not ledger:
                continue
            link = Link(company, ledger, _col(head, r, "Ledger group"), _col(head, r, "Counterparty (Tally company)"),
                        _col(head, r, "Match"), _col(head, r, "Confirmed (Y/N)").upper() in ("Y", "YES"),
                        _col(head, r, "Notes"))
            if link.confirmed or not confirmed_only:
                out.append(link)
    return out


def relationship(a: str, b: str, register: dict[str, Client]) -> str:
    """same entity / same VAT group / same group / different group / unknown."""
    ca, cb = register.get(a), register.get(b)
    if not ca or not cb:
        return "unknown"
    if ca.entity and ca.entity == cb.entity:
        return "same entity"
    if ca.vat_group and ca.vat_group.lower() == cb.vat_group.lower():
        return "same VAT group"
    if ca.group and ca.group.lower() == cb.group.lower():
        return "same group"
    return "different group" if ca.group and cb.group else "unknown"


def check(register: dict[str, Client], links: list[Link] | None = None, open_companies: list[str] | None = None) -> list[str]:
    """Inconsistencies in the register (units of one entity that disagree, unknown names ...)."""
    problems = []
    by_entity: dict[str, list[Client]] = {}
    for c in register.values():
        if c.entity:
            by_entity.setdefault(c.entity, []).append(c)
    for ent, units in sorted(by_entity.items()):
        for attr, label in ENTITY_FIELDS.items():
            values = {getattr(u, attr).strip().lower() for u in units if getattr(u, attr).strip()}
            if len(values) > 1:
                shown = sorted({getattr(u, attr).strip() for u in units if getattr(u, attr).strip()})
                problems.append(f"{ent}: units disagree on {label} ({' / '.join(shown)})")
    for c in register.values():
        if c.vat_period.lower().startswith("q") and not c.vat_quarter_ends:
            problems.append(f"{c.label}: quarterly VAT but no quarter-end months")
    if open_companies is not None:
        missing = [c for c in register if c not in open_companies]
        if missing:
            problems.append("In the register but not open in Tally: " + ", ".join(missing[:10])
                            + (" ..." if len(missing) > 10 else ""))
        new = [c for c in open_companies if c not in register]
        if new and register:
            problems.append("Open in Tally but not in the register: " + ", ".join(new[:10]) + (" ..." if len(new) > 10 else ""))
    for l in links or []:
        if l.company not in register:
            problems.append(f"Inter-company: unknown unit '{l.company}'")
        if l.counterparty and l.counterparty not in register and "outside" not in l.notes.lower():
            problems.append(f"Inter-company: {l.company} / {l.ledger} -> unknown counterparty '{l.counterparty}'")
    return problems


def find_roles(ledgers, groups: dict[str, str]) -> dict[str, list[str]]:
    """Bank / cash / suspense ledgers of one company, from Tally's groups and ledger names."""
    from integrations.tally.client import in_group
    roles = {"bank": [], "cash": [], "suspense": []}
    for l in ledgers:
        if in_group(l.parent, "Bank Accounts", groups) or in_group(l.parent, "Bank OD A/c", groups):
            roles["bank"].append(l.name)
        elif in_group(l.parent, "Cash-in-Hand", groups):
            roles["cash"].append(l.name)
        if "suspense" in l.name.lower() or in_group(l.parent, "Suspense A/c", groups):
            roles["suspense"].append(l.name)
    return roles


# --- suggestions for the draft ---------------------------------------------------------------
def _norm(name: str) -> str:
    n = (name or "").lower()
    n = re.sub(r"^z\s*\(closed\)\s*", "", n)
    n = re.sub(r"\((?:e\s*-?\s*\d+|c)\)", " ", n)
    n = re.sub(r"\b(llc|l\.l\.c|fze|fzco|est|br|branch)\b", " ", n)
    return re.sub(r"[^a-z0-9]+", " ", n).strip()


def suggest_groups(companies: list[str]) -> dict[str, str]:
    """A group per company from the names: 'Mara' for '... Mara ...', else the shared leading words of
    companies with the same first word; then units of one entity share their entity's majority group."""
    first = Counter(_norm(c).split()[0] for c in companies if _norm(c))
    out = {}
    for c in companies:
        words = _norm(c).split()
        if not words:
            continue
        if "mara" in words:
            out[c] = "Mara"
            continue
        if first[words[0]] < 2:
            continue
        same = [_norm(x).split() for x in companies if _norm(x).split()[:1] == words[:1]]
        k = 1
        while all(len(s) > k and s[k] == words[k] for s in same) and k < len(words):
            k += 1
        out[c] = " ".join(words[:k]).title()
    by_entity: dict[str, list[str]] = {}
    for c in companies:
        if entity_code(c):
            by_entity.setdefault(entity_code(c), []).append(c)
    for units in by_entity.values():
        votes = Counter(out[u] for u in units if out.get(u))
        if votes:
            g = votes.most_common(1)[0][0]
            for u in units:
                out[u] = g
    return out


def match_ledger(ledger: str, companies: list[str], own: str) -> tuple[str, str]:
    """(counterparty company, 'exact' / 'partial' / 'close' / '') for an inter-company ledger name."""
    ln = _norm(ledger)
    if not ln:
        return "", ""
    others = [c for c in companies if c != own]
    exact = [c for c in others if _norm(c) == ln]
    if len(exact) == 1:
        return exact[0], "exact"
    part = [c for c in others if _norm(c) and (ln in _norm(c) or _norm(c) in ln)]
    if len(part) == 1:
        return part[0], "partial"
    best, score = "", 0.0
    for c in others:
        s = difflib.SequenceMatcher(None, ln, _norm(c)).ratio()
        if s > score:
            best, score = c, s
    if score >= 0.82:
        return best, "close"
    return "", ""


def draft(tally, store=None, today: date | None = None, existing: dict[str, Client] | None = None,
          existing_links: list[Link] | None = None) -> tuple[str, int]:
    """Build the draft register from Tally. Returns (where it was saved, number of companies)."""
    from openpyxl import Workbook

    from app.brand import BLUE, GREEN, YELLOW, Sheet, save
    from app.storage import free_name
    from integrations.tally.client import in_group
    store = store or _store()
    today = today or date.today()
    existing = existing if existing is not None else load(store)
    old_links = {(l.company, l.ledger): l for l in (existing_links if existing_links is not None
                                                    else load_links(store, confirmed_only=False))}
    companies = sorted(tally.companies(), key=lambda c: c["name"].lower())
    names = [c["name"] for c in companies]
    groups_hint = suggest_groups(names)

    wb = Workbook()
    sh = Sheet(wb, "Units", "Client Register - units, entities and groups",
               f"Drafted by ACE from Tally on {today:%d-%b-%Y} | {len(companies)} units | yellow cells: to be "
               f"completed or confirmed by the team | save as '{REGISTER}' in ACE/tally",
               UNIT_WIDTHS, first=True)
    sh.header(UNIT_COLS)
    link_rows = []
    for c in companies:
        name = c["name"]
        old = existing.get(name)
        roles, last, ledgers, groups = {"bank": [], "cash": [], "suspense": []}, None, [], {}
        try:
            ledgers, groups = tally.ledgers(name), tally.groups(name)
            roles = find_roles(ledgers, groups)
            last = tally.last_voucher_date(name, today - timedelta(days=730), today)
        except Exception as exc:  # noqa: BLE001
            log.warning("register draft %s: %s", name, exc)
        fy = c.get("starting_from")
        ent = old.entity if old and old.entity else entity_code(name)
        g = lambda attr, default="": getattr(old, attr) if old and getattr(old, attr) else default  # noqa: E731
        values = [name, g("short", re.sub(r"\s*\((?:E-\d+|c)\)\s*", " ", name).strip()),
                  ("Y" if old.include else "N") if old else ("N" if re.search(r"\(closed\)", name, re.I) else "Y"),
                  g("group", groups_hint.get(name, "")), ent or (STANDALONE if not old else ""),
                  g("entity_name"), g("trn"), g("vat_period"), g("vat_quarter_ends"), g("ct_trn"),
                  g("fy_start", f"{fy.day}-{fy:%b}" if fy else ""), g("vat_group"), g("ct_group"),
                  g("staff"), g("manager"),
                  "; ".join(old.roles.get("bank") if old and old.roles.get("bank") else roles["bank"]),
                  "; ".join(old.roles.get("cash") if old and old.roles.get("cash") else roles["cash"]),
                  "; ".join(old.roles.get("suspense") if old and old.roles.get("suspense") else roles["suspense"]),
                  f"{last:%d %b %Y}" if last else "none in 2 years", g("notes")]
        todo = {i + 1: YELLOW for i in (5, 6, 7, 13, 14) if not values[i]}       # legal name, TRN, VAT, staff, manager
        if not old:
            todo[4] = YELLOW                                                       # group: suggestion to confirm
        if values[4] == STANDALONE:
            todo[5] = YELLOW
        if str(values[7]).lower().startswith("q") and not values[8]:
            todo[9] = YELLOW
        sh.line(values, fills=todo)

        # inter-company ledgers of this unit
        for l in ledgers:
            grp = next((G for G in IC_GROUPS if in_group(l.parent, G, groups)), None)
            if not grp:
                continue
            prev = old_links.get((name, l.name))
            cp, how = (prev.counterparty, prev.match) if prev and prev.counterparty else match_ledger(l.name, names, name)
            if not cp and grp != "Branch / Divisions" and not prev:
                continue                                   # ordinary customers / suppliers
            link_rows.append((name, l.name, l.parent, cp, how or ("" if cp else "none"),
                              ("Y" if prev.confirmed else "") if prev else ("Y" if how == "exact" else ""),
                              prev.notes if prev else ""))
    sh.note("Group: suggested by ACE from the names - please confirm. Entity code: from the name; 'stand-alone?' = "
            "no code found (leave empty if it is its own entity, or enter its entity code).")
    sh.footer("All clients", "Client Register")

    reg_after = {r[0]: Client(company=r[0]) for r in link_rows}       # relationship needs the units' entity / group
    for row in sh.ws.iter_rows(min_row=sh.header_row + 1, values_only=True):
        if row and row[0] in names:
            e = str(row[4] or "")
            reg_after[row[0]] = Client(company=row[0], group=str(row[3] or ""),
                                       entity="" if e.startswith("stand") else e, vat_group=str(row[11] or ""))
    ic = Sheet(wb, "Inter-company", "Inter-company ledgers - which ledger stands for which unit",
               f"Drafted by ACE from Tally on {today:%d-%b-%Y} | {len(link_rows)} ledgers under Branch / Divisions and "
               "Sundry Debtors / Creditors that look like other units | confirm with Y", IC_WIDTHS, tab=GREEN)
    ic.header(IC_COLS)
    for company, ledger, lgroup, cp, how, conf, notes in link_rows:
        rel = relationship(company, cp, reg_after) if cp else ""
        fills = {}
        if not conf:
            fills[7] = YELLOW
        if not cp:
            fills[4] = YELLOW
        ic.line([company, ledger, lgroup, cp, how, rel, conf, notes], fills=fills)
    if not link_rows:
        ic.note("No inter-company ledgers found.")
    ic.note("Match: exact = same name; partial = one name contains the other; close = similar spelling; none = no unit "
            "found. Only rows with Confirmed = Y are used.")
    ic.footer("All clients", "Client Register")

    h = Sheet(wb, "How to fill", "Client Register - how to fill it in", "Accountability Accountants - ACE",
              [34, 110], tab=BLUE)
    h.header(["Column", "What to put"], freeze=False)
    for row in HELP:
        h.line(list(row))
    h.line(["When done", f"Save the file as '{REGISTER}' in the ACE/tally folder (replace the old one). "
            "ACE reads it before every Tally check and warns about inconsistencies."], style="subtotal")
    h.footer("All clients", "Client Register")
    wb["Units"].sheet_properties.tabColor = GREEN
    name = free_name(store, [], f"Client Register (draft) {today:%Y-%m-%d}.xlsx")
    where = store.save([name], save(wb))
    return where, len(companies)


def match_company(text: str, companies: list[str], register: dict[str, Client]) -> list[str]:
    """Companies a name in a message refers to (Tally name or short name, partial match)."""
    low = text.strip().lower()
    if not low:
        return []
    exact = [c for c in companies if c.lower() == low or (register.get(c) and register[c].short.lower() == low)]
    if exact:
        return exact
    return [c for c in companies if low in c.lower() or (register.get(c) and low in register[c].short.lower())]

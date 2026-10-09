"""Accountability Accountants house style for every document ACE produces (v0.8.1).

Colours come from the firm's logo: navy 'A', green 'a' / ACCOUNTANTS wordmark, blue magnifier,
and the teal table bars on its invoices. Arial everywhere.

Excel layout (Sheet): navy title bar (row 1), grey italic sub-title (row 2), thin green accent
line (row 3), teal column headers with white bold text, green-tint section bands, navy-tint
subtotals, navy top + double bottom border on totals, gridlines off, landscape fit-to-width,
footer "Accountability Accountants - <client> - <engagement> - page/pages".
Yellow fill = to be completed by the team. The firm logo goes top-left on the first sheet when
ACE/branding/logo.png exists on the NAS (never redrawn).
"""
from __future__ import annotations

import io
import logging
import time
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)

NAVY, TEAL, GREEN, BLUE = "002D49", "174D51", "7DB343", "0A8FD0"
NAVY_TINT, GREEN_TINT, GREY = "DCE6EE", "EAF3DC", "595959"
YELLOW = "FFFF00"
SEV_FILL = {"High": "F8CBAD", "Medium": "FFE699", "Low": "EAF3DC"}
AED = '#,##0.00;(#,##0.00);"-"'
AED0 = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
DATE = "dd-mmm-yyyy"
F = "Arial"
FIRM = "Accountability Accountants"
LOGO_PATH = ["branding", "logo.png"]          # in the ACE share


def fill(c: str) -> PatternFill:
    return PatternFill("solid", start_color=c, end_color=c)


_logo_cache: tuple[float, bytes | None] = (0.0, None)


def logo() -> bytes | None:
    """The firm logo (ACE/branding/logo.png on the NAS, or DATA_DIR/branding/logo.png); None if missing."""
    global _logo_cache
    if time.monotonic() - _logo_cache[0] < 600:
        return _logo_cache[1]
    data = None
    try:
        from app.config import settings
        from app.nas import nas
        if nas.configured:
            if nas.exists(settings.ace_share, LOGO_PATH):
                data = nas.read_bytes(settings.ace_share, LOGO_PATH, max_mb=5)
        else:
            p = settings.data_dir.joinpath(*LOGO_PATH)
            data = p.read_bytes() if p.exists() else None
    except Exception as exc:  # noqa: BLE001 - a missing logo must never stop a report
        log.info("logo not available: %s", exc)
    _logo_cache = (time.monotonic(), data)
    return data


class Sheet:
    """One branded worksheet. Rows are written top-down with header(), band(), line(), note()."""

    def __init__(self, wb: Workbook, title: str, heading: str, subtitle: str, widths: list[int], tab=NAVY,
                 first=False, with_logo: bool | None = None, logo_bytes: bytes | None = None):
        self.ws = wb.active if first else wb.create_sheet()
        self.ws.title = title[:31]
        self.ws.sheet_properties.tabColor = tab
        self.ws.sheet_view.showGridLines = False
        for i, w in enumerate(widths, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = w
        n = len(widths)
        self.ncol = n
        o = 0
        img = logo_bytes if logo_bytes is not None else (logo() if (first if with_logo is None else with_logo) else None)
        if img:
            o = self._logo(img)
        top = 1 + o
        self.ws.merge_cells(start_row=top, start_column=1, end_row=top, end_column=max(n, 1))
        c = self.ws.cell(top, 1, heading)
        c.font = Font(name=F, size=14, bold=True, color="FFFFFF")
        c.alignment = Alignment(vertical="center", indent=1)
        for col in range(1, n + 1):
            self.ws.cell(top, col).fill = fill(NAVY)
        self.ws.row_dimensions[top].height = 28
        s = self.ws.cell(top + 1, 1, subtitle)
        s.font = Font(name=F, size=9, italic=True, color=GREY)
        for col in range(1, n + 1):
            self.ws.cell(top + 2, col).fill = fill(GREEN)
        self.ws.row_dimensions[top + 2].height = 4
        self.row = top + 4
        ps = self.ws.page_setup
        ps.orientation = "landscape"
        ps.fitToWidth = 1
        ps.fitToHeight = 0
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True

    def _logo(self, data: bytes) -> int:
        try:
            from openpyxl.drawing.image import Image
            img = Image(io.BytesIO(data))
            ratio = 48 / max(img.height, 1)
            img.height, img.width = 48, int(img.width * ratio)
            self.ws.add_image(img, "A1")
            self.ws.row_dimensions[1].height = 40
            return 1
        except Exception as exc:  # noqa: BLE001
            log.info("logo not placed: %s", exc)
            return 0

    def header(self, labels: list[str], freeze=True):
        for i, lab in enumerate(labels, 1):
            c = self.ws.cell(self.row, i, lab)
            c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
            c.fill = fill(TEAL)
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left" if i == 1 else "center")
        self.ws.row_dimensions[self.row].height = 30
        if freeze:
            self.ws.freeze_panes = self.ws.cell(self.row + 1, 2)
        self.header_row = self.row
        self.row += 1

    def band(self, text: str):
        for col in range(1, self.ncol + 1):
            self.ws.cell(self.row, col).fill = fill(GREEN_TINT)
        c = self.ws.cell(self.row, 1, text)
        c.font = Font(name=F, size=10, bold=True, color=NAVY)
        self.row += 1

    def line(self, values: list[Any], fmts: list[str | None] | None = None, style: str = "line",
             fills: dict[int, str] | None = None):
        bold = style in ("subtotal", "total")
        for i, v in enumerate(values, 1):
            c = self.ws.cell(self.row, i, v)
            c.font = Font(name=F, size=10, bold=bold, color=NAVY if bold else "000000")
            if fmts and i - 1 < len(fmts) and fmts[i - 1]:
                c.number_format = fmts[i - 1]
            if isinstance(v, str) and i > 1:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if style == "subtotal":
                c.fill = fill(NAVY_TINT)
            if style == "total":
                c.border = Border(top=Side("thin", color=NAVY), bottom=Side("double", color=NAVY))
            if fills and i in fills:
                c.fill = fill(fills[i])
        self.row += 1

    def note(self, text: str):
        c = self.ws.cell(self.row, 1, text)
        c.font = Font(name=F, size=9, italic=True, color=GREY)
        self.row += 1

    def footer(self, client: str, engagement: str = "Monthly MIS"):
        self.ws.oddFooter.left.text = f"{FIRM} - {client} - {engagement} - &P/&N"
        self.ws.oddFooter.left.font = "Arial"


def save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
